"""参考文档解析：PDF / DOCX → 文字 + 内嵌图片（带页码）+ 页面渲染图。

页面渲染图供视觉模型理解（doc_vision），嵌入图按"产品页优先"被挑选登上画布。
"""

from dataclasses import dataclass, field
from io import BytesIO


@dataclass
class ParsedDoc:
    text: str = ""
    # (data, mime, page_no 从1开始；docx 无页码概念恒为 0)
    images: list[tuple[bytes, str, int]] = field(default_factory=list)
    # (page_no, png bytes)，仅 PDF 有
    page_renders: list[tuple[int, bytes]] = field(default_factory=list)


_MIN_IMAGE_BYTES = 8 * 1024  # 过滤图标/装饰小图
_MAX_TEXT_CHARS = 20_000
_RENDER_DPI = 72
_MAX_RENDER_PAGES = 16


def parse_pdf(data: bytes) -> ParsedDoc:
    from pypdf import PdfReader

    out = ParsedDoc()
    reader = PdfReader(BytesIO(data))
    texts = []
    for page_no, page in enumerate(reader.pages, start=1):
        try:
            if t := page.extract_text():
                texts.append(t)
        except Exception:
            pass
        try:
            for img in page.images:
                raw = img.data
                if len(raw) >= _MIN_IMAGE_BYTES and (mime := _sniff(raw)):
                    out.images.append((raw, mime, page_no))
        except Exception:
            continue
    out.text = "\n".join(texts)[:_MAX_TEXT_CHARS]

    # 页面渲染（视觉理解用）
    try:
        import fitz  # pymupdf

        doc = fitz.open(stream=data, filetype="pdf")
        for i, page in enumerate(doc):
            if i >= _MAX_RENDER_PAGES:
                break
            pix = page.get_pixmap(dpi=_RENDER_DPI)
            out.page_renders.append((i + 1, pix.tobytes("png")))
        doc.close()
    except Exception:
        pass
    return out


def parse_docx(data: bytes) -> ParsedDoc:
    import docx

    document = docx.Document(BytesIO(data))
    out = ParsedDoc()
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
    out.text = "\n".join(paragraphs)[:_MAX_TEXT_CHARS]

    for rel in document.part.rels.values():
        if "image" in rel.reltype:
            try:
                raw = rel.target_part.blob
                if len(raw) >= _MIN_IMAGE_BYTES and (mime := _sniff(raw)):
                    out.images.append((raw, mime, 0))
            except Exception:
                continue
    return out


def select_images(
    images: list[tuple[bytes, str, int]], product_pages: list[int] | None, limit: int = 8
) -> list[tuple[bytes, str, int]]:
    """挑选登上画布的图：产品页的图优先（页内按尺寸降序），其余补位。"""
    if not product_pages:
        return images[:limit]
    product = sorted(
        [im for im in images if im[2] in product_pages],
        key=lambda im: (product_pages.index(im[2]), -len(im[0])),
    )
    rest = [im for im in images if im[2] not in product_pages]
    return (product + rest)[:limit]


def _sniff(data: bytes) -> str | None:
    if data[:4] == b"\x89PNG":
        return "image/png"
    if data[:3] == b"\xff\xd8\xff":
        return "image/jpeg"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp"
    return None


def parse(filename: str, data: bytes) -> ParsedDoc:
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        return parse_pdf(data)
    if ext in ("docx", "doc"):
        return parse_docx(data)  # .doc 旧格式大概率失败，由调用方兜底报错
    raise ValueError(f"unsupported document type: {ext}")
